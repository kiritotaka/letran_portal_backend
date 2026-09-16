"""Read product identities only; never evaluate Excel formulas or infer codes."""
import csv
from io import BytesIO, StringIO
import unicodedata

from openpyxl import load_workbook

from app.services.file_policies import POLICIES
from app.services.file_validation import validate_file


class ImportFailure(Exception):
    def __init__(self, code):
        self.code = code


def text(value):
    return unicodedata.normalize('NFC', value).strip() if isinstance(value, str) else value


def header_positions(row):
    values = [text(value) for value, _ in row]
    headers = ('Mã TP/BTP', 'Tên TP/BTP')
    if all(values.count(h) == 1 for h in headers):
        return tuple(values.index(h) for h in headers)


def parse_rows(rows, sheet_name, checkpoint):
    products, errors = {}, []
    counts = {'source_rows': 0, 'duplicate_rows': 0, 'ignored_rows': 0}
    positions = None
    for number, row in enumerate(rows, 1):
        if number > 100000:
            raise ImportFailure('IMPORT_ROW_LIMIT')
        if number % 500 == 0:
            checkpoint()
        if positions is None:
            positions = header_positions(row)
            if positions is None and number >= 20:
                raise ImportFailure('IMPORT_HEADERS_NOT_FOUND')
            continue
        selected = [row[i] if i < len(row) else (None, None) for i in positions]
        code, name = [text(cell[0]) for cell in selected]
        # Warehouse reports contain a final numeric total outside these columns.
        if code in (None, '') and name in (None, ''):
            counts['ignored_rows'] += 1
            continue
        counts['source_rows'] += 1
        reason = None
        if any(kind in ('f', 'e') for _, kind in selected):
            reason = 'FORMULA_OR_CELL_ERROR'
        elif not isinstance(code, str) or not code or len(code) > 100:
            reason = 'INVALID_PRODUCT_CODE'
        elif not isinstance(name, str) or not name or len(name) > 500:
            reason = 'INVALID_PRODUCT_NAME'
        elif any(ord(c) < 32 for c in code + name):
            reason = 'INVALID_CHARACTERS'
        if reason:
            errors.append({'source_row': number, 'product_code': str(code or '')[:100],
                           'product_name': str(name or '')[:500], 'error_code': reason})
        elif code in products:
            original = products[code]
            original['occurrences'] += 1
            if original['product_name'] != name:
                original['error_code'] = 'CONFLICTING_PRODUCT_NAME'
                errors.append({'source_row': number, 'product_code': code,
                               'product_name': name, 'error_code': 'CONFLICTING_PRODUCT_NAME'})
            else:
                counts['duplicate_rows'] += 1
        else:
            products[code] = {'source_row': number, 'product_code': code,
                              'product_name': name, 'occurrences': 1, 'error_code': None}
        if len(products) + len(errors) > 10000:
            raise ImportFailure('IMPORT_PREVIEW_LIMIT')
    if positions is None:
        raise ImportFailure('IMPORT_HEADERS_NOT_FOUND')
    if not products and not errors:
        raise ImportFailure('IMPORT_EMPTY')
    return {'rows': list(products.values()) + errors,
            'summary': {**counts, 'unique_codes': len(products), 'sheet_name': sheet_name}}


def parse_products(filename, content, sheet_name=None, checkpoint=lambda: None):
    validate_file(filename, content, POLICIES['product_import'])
    if filename.lower().endswith('.csv'):
        if sheet_name:
            raise ImportFailure('IMPORT_SHEET_NOT_FOUND')
        reader = csv.reader(StringIO(content.decode('utf-8-sig')), strict=True)
        return parse_rows(([(v, 's') for v in row] for row in reader), None, checkpoint)
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ImportFailure('IMPORT_SHEET_NOT_FOUND')
            sheet = workbook[sheet_name]
        elif len(workbook.worksheets) == 1:
            sheet = workbook.worksheets[0]
        else:
            raise ImportFailure('IMPORT_SHEET_REQUIRED')
        sheet.reset_dimensions()
        rows = ([(c.value, c.data_type) for c in row]
                for row in sheet.iter_rows(max_col=100))
        return parse_rows(rows, sheet.title, checkpoint)
    finally:
        workbook.close()
